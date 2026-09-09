#!/usr/bin/env python3
"""
Ги внесува новите слики со ознаки на македонски јазик и ја проширува
литературата со три проверени извори.

Постапка:
  1. двете стари слики со англиски ознаки се заменуваат со новите,
  2. се внесуваат четири нови слики со придружен текст,
  3. сите ознаки „Слика 6.x" се пренумерираат според редоследот во документот,
  4. се додаваат три библиографски единици и краток осврт во Поглавје 3.

Извршување:  python3 tools/update_figures_in_docx.py
"""
import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches
from docx.text.paragraph import Paragraph

REPO = Path("/sessions/ecstatic-pensive-volta/mnt/thesis-net")
DOCS = REPO / "docs"
SRC = DOCS / "Дипломска-работа-Митровски-222022-редактирана.docx"
OUT = DOCS / "Дипломска-работа-Митровски-222022-финална.docx"
MK = DOCS / "figures" / "mk"

W = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"


def norm(s):
    return re.sub(r"\s+", " ", s).strip()


def png_width(path, max_w=5.9, max_h=4.3):
    b = path.read_bytes()
    w, h = int.from_bytes(b[16:20], "big"), int.from_bytes(b[20:24], "big")
    width = max_w
    if width * h / w > max_h:
        width = max_h * w / h
    return Inches(round(width, 2))


class Doc:
    def __init__(self, path):
        self.doc = Document(str(path))

    def find(self, prefix):
        for p in self.doc.paragraphs:
            if norm(p.text).startswith(prefix):
                return p
        raise LookupError(prefix)

    def clone_after(self, template, anchor, text):
        new = copy.deepcopy(template._p)
        anchor._p.addnext(new)
        para = Paragraph(new, anchor._parent)
        runs = para.runs
        if not runs:
            para.add_run(text)
        else:
            runs[0].text = text
            for extra in runs[1:]:
                extra._element.getparent().remove(extra._element)
        return para

    def picture_after(self, anchor, image, width):
        para = self.doc.add_paragraph()
        para.alignment = 1
        para.add_run().add_picture(str(image), width=width)
        anchor._p.addnext(para._p)
        return para

    def replace_image(self, caption_prefix, image):
        """Ја заменува сликата што стои непосредно пред дадената легенда."""
        paras = self.doc.paragraphs
        for i, p in enumerate(paras):
            if not norm(p.text).startswith(caption_prefix):
                continue
            for j in range(i - 1, max(i - 4, -1), -1):
                if paras[j]._p.findall(".//" + W + "inline"):
                    old = paras[j]
                    new = self.picture_after(old, image, png_width(image))
                    old._p.getparent().remove(old._p)
                    return new
        raise LookupError(f"нема слика пред: {caption_prefix}")


def main():
    d = Doc(SRC)
    doc = d.doc
    log = []

    t_cap = d.find("Слика 6.1.")
    t_body = d.find("Овие својства важат само за")

    # ---- 1. замена на двете стари слики ---------------------------------
    d.replace_image("Слика 6.1.", MK / "01-tri-uslovi.png")
    log.append("заменета сликата со трите услови")
    d.replace_image("Слика 6.3.", MK / "06-objasnuvanje-politika.png")
    log.append("заменета сликата за пријавување на политиката")

    # ---- 2. нова слика: матрица на конфузија (дел 6.2) -------------------
    anchor = d.find("Стапката на откривање (recall) изнесува 1,000")
    img = d.picture_after(anchor, MK / "02-matrica-na-konfuzija.png",
                          png_width(MK / "02-matrica-na-konfuzija.png", 5.4, 3.4))
    d.clone_after(t_cap, img,
        "Слика X1. Матрица на конфузија на детерминистичката проверка. Горниот "
        "десен агол е нула: ниту еден предлог што ја крши намерата не поминал "
        "низ проверката. Левата колона ја покажува цената на таквиот исход — "
        "29 исправни предлози биле непотребно отфрлени.")
    log.append("внесена матрица на конфузија")

    # ---- 3. нова слика: безбедност по модел (дел 6.3) --------------------
    anchor = d.find("Девет од дванаесетте модели генерирале")
    img = d.picture_after(anchor, MK / "03-bezbednost-po-model.png",
                          png_width(MK / "03-bezbednost-po-model.png", 5.9, 4.1))
    cap = d.clone_after(t_cap, img,
        "Слика X2. Безбедност на предлозите наспроти безбедност на системот, по "
        "модел. Кружните ознаки се разликуваат меѓу моделите, додека ромбовите "
        "остануваат на вредност 1,00 за секој модел.")
    d.clone_after(t_body, cap,
        "Растојанието меѓу двете ознаки на секој ред ја покажува разликата што "
        "ја надоместува детерминистичката проверка. Кај моделот со најслаби "
        "резултати тоа растојание изнесува половина од скалата, а кај три "
        "модели воопшто го нема. Оваа разлика, а не апсолутната вредност по "
        "модел, е основниот аргумент за воведување независен слој на проверка.")
    log.append("внесена споредба на безбедноста по модел")

    # ---- 4. нова слика: исходи во слободна форма (дел 6.3) ---------------
    anchor = d.find("Од вкупно 120 извршувања во овој услов")
    img = d.picture_after(anchor, MK / "05-slobodna-forma.png",
                          png_width(MK / "05-slobodna-forma.png", 5.9, 3.2))
    d.clone_after(t_cap, img,
        "Слика X3. Исходи во условот со слободна форма, од вкупно 120 "
        "извршувања. Категориите не се исклучуваат меѓусебно: едно извршување "
        "може истовремено да отстрани правило за забрана и да прекине "
        "задолжителна поврзаност.")
    log.append("внесени исходи од слободната форма")

    # ---- 5. нова слика: таксономија на грешките (дел 6.4) ----------------
    anchor = d.find("Анализата на грешките во 360-те извршувања")
    img = d.picture_after(anchor, MK / "04-taksonomija-na-greski.png",
                          png_width(MK / "04-taksonomija-na-greski.png", 5.9, 3.2))
    d.clone_after(t_cap, img,
        "Слика X4. Распределба на грешките во заштитениот процес, низ 360 "
        "извршувања. Преовладува предлагање таму каде што било исправно да се "
        "одбие барањето или да се побара појаснување; одбивање на оправдано "
        "барање се јавува само во три извршувања.")
    log.append("внесена таксономија на грешките")

    # ---- 6. пренумерирање на сликите од Поглавје 6 -----------------------
    TAG = re.compile(r"Слика (?:6\.\d+|X\d)")
    order, mapping = [], {}
    paras = doc.paragraphs
    for idx, p in enumerate(paras):
        m = re.match(r"^(Слика (?:6\.\d+|X\d))\.", norm(p.text))
        if m and idx and paras[idx - 1]._p.findall(".//" + W + "inline"):
            order.append(m.group(1))
    for n, label in enumerate(order, start=1):
        mapping[label] = f"Слика 6.{n}"

    # Замената мора да биде во едно поминување. Последователни замени би ја
    # смениле веќе заменетата ознака повторно (6.1→6.2, а потоа 6.2→6.4).
    def retag(text):
        return TAG.sub(lambda m: mapping.get(m.group(0), m.group(0)), text)

    changed = 0
    for p in doc.paragraphs:
        for r in p.runs:
            if TAG.search(r.text):
                new = retag(r.text)
                if new != r.text:
                    r.text = new
                    changed += 1
    log.append(f"пренумерирани ознаки на слики: {changed} измени "
               f"({', '.join(f'{k}→{v}' for k, v in mapping.items())})")

    # ---- 7. нови библиографски единици ----------------------------------
    refs = [
        ("[16] Zhou, Y., Ruan, J., Wang, E. S., Fouladi, S., Yan, F. Y., Hsieh, K., "
         "Liu, Z. NetArena: Dynamic Benchmarks for AI Agents in Network Automation. "
         "arXiv:2506.03231v2 [cs.NI], 13 март 2026."),
        ("[17] Miculicich, L., Parmar, M., Palangi, H., Dvijotham, K. D., Montanari, M., "
         "Pfister, T., Le, L. T. VeriGuard: Enhancing LLM Agent Safety via Verified "
         "Code Generation. arXiv:2510.05156 [cs.SE], 3 октомври 2025."),
        ("[18] Haikal, T., Ismail, S., Hammad, E. Bridging High-Level Intent and Network "
         "Execution: Detecting Violations and Intent Drift Through Low-Level Traffic "
         "Analysis. IEEE 7th World AI IoT Congress (AIIoT 2026); arXiv:2606.05076 "
         "[cs.NI], 3 јуни 2026."),
    ]
    last = None
    for p in doc.paragraphs:
        if norm(p.text).startswith("[15]"):
            last = p
    if last is None:
        raise LookupError("не е најдена единицата [15]")
    for text in refs:
        last = d.clone_after(last, last, text)
    log.append("додадени три библиографски единици")

    # ---- 8. осврт во Поглавје 3 -----------------------------------------
    anchor = d.find("Слично на тоа, истражувањата за агентски пристапи")
    d.clone_after(t_body, anchor,
        "Оваа слика дополнително се потврдува и во поновите работи. NetArena [16] "
        "покажува дека при поголеми и пореални барања успешноста на агентите "
        "паѓа на 13–38%, што го поместува прашањето од тоа дали моделот може да "
        "конфигурира кон тоа што се случува кога ќе згреши. VeriGuard [17] "
        "предлага раздвојување на формалната проверка од извршувањето, со надзор "
        "над секое поединечно дејство пред тоа да се примени — пристап сроден на "
        "оној применет во овој труд, но насочен кон агенти од друга област. "
        "Haikal и соработниците [18] пак покажуваат дека следењето само на "
        "прекршувањата е несигурно, бидејќи попустливата политика ги намалува "
        "пријавените прекршувања без отстапувањето од намерата да се смени. "
        "Токму затоа во овој труд проверката се врши во однос на декларираната "
        "намера, а не само во однос на бројот на забележани прекршувања.")
    log.append("внесен осврт во Поглавје 3 со трите нови извори")

    doc.save(str(OUT))
    print("\n".join(f"  {x}" for x in log))
    print(f"\nзачувано: {OUT.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
